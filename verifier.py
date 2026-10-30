import math
import re
import subprocess
import sys
import time
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent
TSP_BINARY = ROOT / "cpp" / "build" / "tsp"
CHECKER_PATH = ROOT / "Solution_checker.xlsx"
VIS_DIR = ROOT / "cpp" / "vis"
INSTANCES = ["TSPA", "TSPB"]
ALGORITHMS = [
    "random",
    "nearest-neighbor-end",
    "nearest-neighbor-anywhere",
    "greedy-cycle",
    "2-regret-r1-d1-nearest-neighbor-anywhere",
    "2-regret-r1-d1-greedy-cycle",
    "2-regret-r1-d0-nearest-neighbor-anywhere",
    "2-regret-r1-d0-greedy-cycle",
    "2-regret-r2-d1-nearest-neighbor-anywhere",
    "2-regret-r2-d1-greedy-cycle",
    "2-regret-r1-d2-nearest-neighbor-anywhere",
    "2-regret-r1-d2-greedy-cycle",
    # lab#3
    "ls-steepest-two-edges-random",
    "ls-steepest-two-nodes-random",
    "ls-steepest-two-edges-greedy",
    "ls-steepest-two-nodes-greedy",
    "ls-greedy-two-edges-random",
    "ls-greedy-two-nodes-random",
    "ls-greedy-two-edges-greedy",
    "ls-greedy-two-nodes-greedy",

]

RESULT_RE = re.compile(
    r"cost:\s*(-?\d+(?:\.\d+)?),\s*time:\s*([0-9]+(?:\.[0-9]+)?)\s*s"
)
ATTEMPT_TIMES_RE = re.compile(
    r"attempt_times:\s*([0-9.eE+-]+(?:\s+[0-9.eE+-]+)*)"
)
READABLE_NAMES = {
    "2-regret-r1-d1-nearest-neighbor-anywhere": "2-regret (regret=1, delta=1) nearest-neighbor-anywhere",
    "2-regret-r1-d1-greedy-cycle": "2-regret (regret=1, delta=1) greedy-cycle",
    "2-regret-r1-d0-nearest-neighbor-anywhere": "2-regret (regret=1, delta=0) nearest-neighbor-anywhere",
    "2-regret-r1-d0-greedy-cycle": "2-regret (regret=1, delta=0) greedy-cycle",
    "2-regret-r2-d1-nearest-neighbor-anywhere": "2-regret (regret=2, delta=1) nearest-neighbor-anywhere",
    "2-regret-r2-d1-greedy-cycle": "2-regret (regret=2, delta=1) greedy-cycle",
    "2-regret-r1-d2-nearest-neighbor-anywhere": "2-regret (regret=1, delta=2) nearest-neighbor-anywhere",
    "2-regret-r1-d2-greedy-cycle": "2-regret (regret=1, delta=2) greedy-cycle",
    "ls-steepest-two-edges-random": "LS steepest, two-edges, random start",
    "ls-steepest-two-nodes-random": "LS steepest, two-nodes, random start",
    "ls-steepest-two-edges-greedy": "LS steepest, two-edges, greedy start",
    "ls-steepest-two-nodes-greedy": "LS steepest, two-nodes, greedy start",
    "ls-greedy-two-edges-random": "LS greedy, two-edges, random start",
    "ls-greedy-two-nodes-random": "LS greedy, two-nodes, random start",
    "ls-greedy-two-edges-greedy": "LS greedy, two-edges, greedy start",
    "ls-greedy-two-nodes-greedy": "LS greedy, two-nodes, greedy start",
}


def load_instance(path):
    data = np.loadtxt(path, delimiter=";", dtype=np.int64)
    coords = data[:, :2]
    costs = data[:, 2]
    diff = coords[:, None, :] - coords[None, :, :]
    distance_matrix = np.floor(np.sqrt((diff**2).sum(axis=2)) + 0.5).astype(np.int64)
    return coords, costs, distance_matrix


def check_solution(cycle, n, target_count, distance_matrix, costs):
    if len(cycle) != target_count:
        raise ValueError(f"{len(cycle)} nodes, expected {target_count}")
    visited = set()
    for i, node in enumerate(cycle):
        if not 0 <= node < n:
            raise ValueError(f"node {i} out of range: {node}")
        if node in visited:
            raise ValueError(f"node {i} is a duplicate: {node}")
        visited.add(node)
    total = 0
    for i in range(len(cycle)):
        total += int(distance_matrix[cycle[i]][cycle[(i + 1) % len(cycle)]]) + int(
            costs[cycle[i]]
        )
    return total


CHECKER_TABLES = {}


def checker_objective(checker_path, sheet, cycle):
    key = (str(checker_path), sheet)
    if key not in CHECKER_TABLES:
        ws = openpyxl.load_workbook(checker_path)[sheet]
        CHECKER_TABLES[key] = {
            int(ws.cell(r, 1).value): (
                ws.cell(r, 2).value,
                ws.cell(r, 3).value,
                ws.cell(r, 4).value,
            )
            for r in range(3, 203)
        }
    table = CHECKER_TABLES[key]
    total = 0
    for i, node in enumerate(cycle):
        x1, y1, c = table[node]
        x2, y2, _ = table[cycle[(i + 1) % len(cycle)]]
        total += int(math.sqrt((x1 - x2) ** 2 + (y1 - y2) ** 2) + 0.5) + c
    return total


def visualize(coords, costs, cycle, title, path):
    fig, ax = plt.subplots(figsize=(6, 5))
    sc = ax.scatter(
        coords[:, 0],
        coords[:, 1],
        c=costs,
        s=30,
        cmap="viridis",
        edgecolors="black",
        linewidth=0.4,
    )
    pts = coords[cycle + [cycle[0]]]
    ax.plot(pts[:, 0], pts[:, 1], color="red", linewidth=1.2)
    ax.scatter(
        coords[cycle, 0],
        coords[cycle, 1],
        s=80,
        marker="s",
        facecolors="none",
        edgecolors="red",
        linewidth=1,
    )
    ax.set_title(title)
    ax.set_xlabel("X Coordinate")
    ax.set_ylabel("Y Coordinate")
    fig.colorbar(sc, ax=ax, label="Node Cost")
    fig.tight_layout()
    fig.savefig(path, dpi=150)
    plt.close(fig)


def verify(name, instance):
    coords, costs, distance_matrix = load_instance(ROOT / f"{instance}.csv")
    n = len(costs)
    target_count = math.ceil(0.5 * n)

    try:
        process = subprocess.run(
            [str(TSP_BINARY), name],
            input=(ROOT / f"{instance}.csv").read_text(),
            capture_output=True,
            text=True,
        )
    except OSError as error:
        return f"wrong {name} {instance}: {error}"

    if process.returncode != 0:
        return f"wrong {name} {instance}: exit {process.returncode}: {process.stderr.strip()}"

    attempts = 200 if name == "random" else n
    lines = process.stdout.strip().splitlines()
    if len(lines) != attempts:
        return (
            f"wrong {name} {instance}: expected {attempts} solutions, got {len(lines)}"
        )

    tours = []
    values = []
    for index, line in enumerate(lines):
        try:
            tour = [int(value) for value in line.split(",")]
        except ValueError:
            return f"wrong {name} {instance}: malformed solution {index}"
        try:
            value = check_solution(tour, n, target_count, distance_matrix, costs)
        except ValueError as error:
            return f"wrong {name} {instance}: solution {index}: {error}"
        tours.append(tour)
        values.append(value)

    min_index = min(range(len(values)), key=lambda i: values[i])
    max_index = max(range(len(values)), key=lambda i: values[i])
    mean = sum(values) / len(values)

    match = RESULT_RE.search(process.stderr)
    if not match:
        return f"wrong {name} {instance}: malformed score/time"
    reported = float(match.group(1))

    times_match = ATTEMPT_TIMES_RE.search(process.stderr)
    if not times_match:
        return f"wrong {name} {instance}: malformed attempt_times"
    times = [float(x) for x in times_match.group(1).split()]
    if len(times) != attempts:
        return (
            f"wrong {name} {instance}: expected {attempts} attempt times, got {len(times)}"
        )

    if reported != values[min_index]:
        return (
            f"wrong {name} {instance}: score {reported:g}, expected {values[min_index]}"
        )

    for index in {min_index, max_index}:
        if checker_objective(CHECKER_PATH, instance, tours[index]) != values[index]:
            return f"wrong {name} {instance}: checker objective mismatch for solution {index}"

    VIS_DIR.mkdir(exist_ok=True)
    label = READABLE_NAMES.get(name, name)
    visualize(
        coords,
        costs,
        tours[min_index],
        f"{label} {instance} min (objective {values[min_index]})",
        VIS_DIR / f"{name}_{instance}_min.png",
    )
    visualize(
        coords,
        costs,
        tours[max_index],
        f"{label} {instance} max (objective {values[max_index]})",
        VIS_DIR / f"{name}_{instance}_max.png",
    )

    return {"values": values, "times": times}


def objective_cell(values):
    return f"{sum(values)/len(values):.2f} ({min(values)} – {max(values)})"


def time_cell(times):
    return (
        f"{sum(times)/len(times):.4f} ({min(times):.4f} – {max(times):.4f})"
    )


def main():
    names = sys.argv[1:] or ALGORITHMS
    for name in names:
        if name not in ALGORITHMS:
            print(f"wrong unknown algorithm: {name}", file=sys.stderr)
            return 1
    started = time.perf_counter()
    results = {instance: {} for instance in INSTANCES}
    for name in names:
        for instance in INSTANCES:
            result = verify(name, instance)
            if isinstance(result, str):
                print(result, file=sys.stderr)
                return 1
            results[instance][name] = result

    print("### objective")
    print()
    print("|method|" + "|".join(INSTANCES) + "|")
    print("|---|" + "|".join(["---:" for _ in INSTANCES]) + "|")
    for name in names:
        row = [READABLE_NAMES.get(name, name)]
        for instance in INSTANCES:
            row.append(objective_cell(results[instance][name]["values"]))
        print("|" + "|".join(row) + "|")
    print()

    print("### time")
    print()
    print("|method|" + "|".join(INSTANCES) + "|")
    print("|---|" + "|".join(["---:" for _ in INSTANCES]) + "|")
    for name in names:
        row = [READABLE_NAMES.get(name, name)]
        for instance in INSTANCES:
            row.append(time_cell(results[instance][name]["times"]))
        print("|" + "|".join(row) + "|")
    print()

    print(f"total: {time.perf_counter() - started:.3f} s")
    print("ok", file=sys.stderr)
    return 0


if __name__ == "__main__":
    sys.exit(main())
