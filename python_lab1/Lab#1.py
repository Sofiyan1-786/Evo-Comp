import sys
import math
import random
import numpy as np
import openpyxl
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt

def load_instance(path):
    data = np.loadtxt(path, delimiter=';', dtype=np.int64)
    coords = data[:, :2]
    costs = data[:, 2]
    # Precompute full distance matrix right after reading the instance (rounded mathematically)
    diff = coords[:, None, :] - coords[None, :, :]
    distance_matrix = np.floor(np.sqrt((diff ** 2).sum(axis=2)) + 0.5).astype(np.int64)
    return coords, costs, distance_matrix


def objective(cycle, distance_matrix, costs):
    #idx here is the array of indices of the cycle, 
    #nxt is the array of indices of the next node in the cycle (rolled by -1)
    idx = np.asarray(cycle)
    nxt = np.roll(idx, -1)
    return int(distance_matrix[idx, nxt].sum() + costs[idx].sum())


def check_solution(cycle, n, target_count, distance_matrix, costs):

    """Independent validity check: size, uniqueness, index range, recomputed objective."""
    assert len(cycle) == target_count, "wrong number of nodes"
    assert len(set(cycle)) == len(cycle), "duplicate nodes"
    assert all(0 <= v < n for v in cycle), "index out of range"
    total = 0
    for i in range(len(cycle)):
        total += int(distance_matrix[cycle[i]][cycle[(i + 1) % len(cycle)]]) + int(costs[cycle[i]])
    return total


def random_solution_baseline(distance_matrix, costs, target_count):
    """Return a random valid cycle of exactly target_count nodes and its objective."""
    n = len(costs)
    cycle = random.sample(range(n), target_count)
    return cycle, objective(cycle, distance_matrix, costs)


#Nearest neighbor considering adding the node only at the end of the current path
def nearest_neighbor_end(start_node, distance_matrix, costs, target_count):
    n = len(costs)
    cycle_nn_end = [start_node]
    visited = np.zeros(n, dtype=bool)
    visited[start_node] = True

    while len(cycle_nn_end) < target_count:
        last_node = cycle_nn_end[-1]
        # The local step cost change of adding 'neighbor' to the end of the path
        delta = (distance_matrix[last_node] + costs).astype(float)
        delta[visited] = np.inf
        best_node = int(np.argmin(delta))
        cycle_nn_end.append(best_node)
        visited[best_node] = True

    # Calculate final full cycle objective value (including closing edge back to start)
    return cycle_nn_end, objective(cycle_nn_end, distance_matrix, costs)


#Nearest neighbor considering adding the node at all possible position, i.e. at the end, at the beginning, or at any place inside the current path
def nearest_neighbor_solution_all_positions(start_node, distance_matrix, costs, target_count):
    n = len(costs)
    cycle_nn_all = [start_node]
    visited = np.zeros(n, dtype=bool)
    visited[start_node] = True

    # Step 1: Find the best second node to create an initial path sequence
    delta = (distance_matrix[start_node] + costs).astype(float)
    delta[visited] = np.inf
    best_second_node = int(np.argmin(delta))
    cycle_nn_all.append(best_second_node)
    visited[best_second_node] = True

    # Step 2: Now that we have a valid sequence of 2 nodes, build up to target_count
    while len(cycle_nn_all) < target_count:
        path = np.array(cycle_nn_all)
        cand = np.flatnonzero(~visited)
        c = costs[cand]

        # rows: candidates, columns: insertion positions 0..len(path)
        front = distance_matrix[cand, path[0]] + c
        end = distance_matrix[path[-1], cand] + c
        inner = (distance_matrix[cand][:, path[:-1]]
                 + distance_matrix[cand][:, path[1:]]
                 - distance_matrix[path[:-1], path[1:]][None, :]
                 + c[:, None])
        deltas = np.concatenate([front[:, None], inner, end[:, None]], axis=1)

        flat = int(np.argmin(deltas))
        ci, best_position = divmod(flat, deltas.shape[1])
        best_node = int(cand[ci])

        cycle_nn_all.insert(best_position, best_node)
        visited[best_node] = True

    return cycle_nn_all, objective(cycle_nn_all, distance_matrix, costs)


#now greedy cycle
def greedy_cycle(start_node, distance_matrix, costs, target_count):
    n = len(costs)
    # 1. Initialize with the start node
    cycle_greedy = [start_node]
    visited = np.zeros(n, dtype=bool)
    visited[start_node] = True

    # 2. To form an initial loop, find the absolute best secondary node
    # Multiplied by 2 because it forms a 2-way edge loop (start <-> neighbor)
    delta = (distance_matrix[start_node] * 2 + costs).astype(float)
    delta[visited] = np.inf
    best_second_node = int(np.argmin(delta))
    cycle_greedy.append(best_second_node)
    visited[best_second_node] = True

    # 3. Main loop: Expand the cycle until target_count is reached
    while len(cycle_greedy) < target_count:
        cyc = np.array(cycle_greedy)
        nxt = np.roll(cyc, -1)
        cand = np.flatnonzero(~visited)

        # Check EVERY edge currently in the cycle, for every unvisited node
        deltas = (distance_matrix[cand][:, cyc]
                  + distance_matrix[cand][:, nxt]
                  - distance_matrix[cyc, nxt][None, :]
                  + costs[cand][:, None])

        flat = int(np.argmin(deltas))
        ci, i = divmod(flat, deltas.shape[1])
        best_node = int(cand[ci])

        # Insert the best node into its most optimal edge slot (right after node u)
        cycle_greedy.insert(i + 1, best_node)
        visited[best_node] = True

    return cycle_greedy, objective(cycle_greedy, distance_matrix, costs)


def run_instance(name, path):
    coords, costs, distance_matrix = load_instance(path)
    n = len(costs)
    target_count = math.ceil(0.5 * n)  # Exactly 50% coverage required (rounded up for odd n)

    methods = {
        "Random": [random_solution_baseline(distance_matrix, costs, target_count)
                   for _ in range(200)],
        # Deterministic methods: one solution per starting node
        "Nearest Neighbor End": [nearest_neighbor_end(s, distance_matrix, costs, target_count)
                                 for s in range(n)],
        "Nearest Neighbor Anywhere": [nearest_neighbor_solution_all_positions(s, distance_matrix, costs, target_count)
                                      for s in range(n)],
        "Greedy Cycle": [greedy_cycle(s, distance_matrix, costs, target_count)
                         for s in range(n)],
    }

    best = {}
    print(f"\n {name} (n={n}, selected={target_count})")
    for method, sols in methods.items():
        objectives = [o for _, o in sols]
        cycle, obj = min(sols, key=lambda item: item[1])
        checked = check_solution(cycle, n, target_count, distance_matrix, costs)
        assert checked == obj
        best[method] = (cycle, obj)
        print(f"{method}: min={min(objectives)} max={max(objectives)} avg={sum(objectives) / len(objectives):.2f}")
        print(f"  Best Solution (0-based indices): {cycle}")

    fig, axes = plt.subplots(2, 2, figsize=(16, 14))
    for ax, (method, (cycle, obj)) in zip(axes.ravel(), best.items()):
        sc = ax.scatter(coords[:, 0], coords[:, 1], c=costs, s=30, cmap='viridis',
                        edgecolors='black', linewidth=0.4)
        pts = coords[cycle + [cycle[0]]]
        ax.plot(pts[:, 0], pts[:, 1], color='red', linewidth=1.2)
        ax.scatter(coords[cycle, 0], coords[cycle, 1], s=80, marker='s',
                   facecolors='none', edgecolors='red', linewidth=1)
        ax.set_title(f"{name} - {method} (objective {obj})")
        ax.set_xlabel("X Coordinate")
        ax.set_ylabel("Y Coordinate")
        fig.colorbar(sc, ax=ax, label='Node Cost')
    fig.tight_layout()
    fig.savefig(f"{name}_best.png", dpi=150)
    plt.close(fig)

    return best

def checker_objective(checker_path, sheet, cycle):
    # Re-implements the Solution_checker.xlsx formulas: LOOKUP by node id, INT(SQRT(dx^2+dy^2)+0.5), closing edge F103 = F3
    ws = openpyxl.load_workbook(checker_path)[sheet]
    table = {int(ws.cell(r, 1).value): (ws.cell(r, 2).value, ws.cell(r, 3).value, ws.cell(r, 4).value)
             for r in range(3, 203)}
    total = 0
    for i, node in enumerate(cycle):
        x1, y1, c = table[node]
        x2, y2, _ = table[cycle[(i + 1) % len(cycle)]]
        total += int(math.sqrt((x1 - x2) ** 2 + (y1 - y2) ** 2) + 0.5) + c
    return total


def fill_checker(checker_path, out_path, solutions):
    # solutions: {sheet_name: cycle}; writes node ids into F3:F102 so Excel can recompute K2
    wb = openpyxl.load_workbook(checker_path)
    for sheet, cycle in solutions.items():
        ws = wb[sheet]
        for i, node in enumerate(cycle):
            ws.cell(3 + i, 6).value = node
    wb.save(out_path)


def verify_with_checker(results, checker_path="Solution_checker.xlsx"):
    # results: {instance_name: {method: (cycle, obj)}}
    methods = next(iter(results.values())).keys()
    for method in methods:
        tag = method.replace(" ", "_")
        fill_checker(checker_path, f"Solution_checker_{tag}.xlsx",
                     {name: results[name][method][0] for name in results})
        for name in results:
            cycle, obj = results[name][method]
            chk = checker_objective(checker_path, name, cycle)
            print(f"{name} | {method}: code={obj} checker={chk} -> {'OK' if chk == obj else 'MISMATCH'}")


if __name__ == "__main__":
    random.seed(123)
    paths = ["TSPA.csv", "TSPB.csv"]
    results = {}
    for p in paths:
        name = p.rsplit("/", 1)[-1].rsplit(".", 1)[0]
        results[name] = run_instance(name, p)
    verify_with_checker(results)